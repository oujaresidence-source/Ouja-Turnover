# -*- coding: utf-8 -*-
"""
permits.importer — nothing dropped, no ID stored, no date guessed in silence.

Run: python3 -m unittest tests.test_permits_importer
"""

import io
import json
import os
import shutil
import sqlite3
import sys
import tempfile
import unittest
from contextlib import closing

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from brain import db as bdb                       # noqa: E402
from permits import db, engine, importer          # noqa: E402

FAKE_ID = "1" + "0" * 5 + "4321"                   # national-ID-shaped, built so no literal ID sits in the repo


def csv_bytes(rows, enc="utf-8-sig"):
    return "\n".join(",".join(r) for r in rows).encode(enc)


def xlsx_bytes(rows, header_at=3, link=None):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="شركة عوجا — قائمة التصاريح")
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ws.cell(row=header_at + i, column=j + 1, value=v)
            if link and (i, j) == link[0]:
                c.hyperlink = link[1]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class DbCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="permits_imp_")
        bdb.set_db_path_for_tests(os.path.join(self.tmp, "brain.db"))
        db.reset_init_cache()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)


class TestReading(DbCase):

    HEAD = ["اسم الوحدة", "رقم التصريح", "تاريخ الانتهاء", "رقم الهوية", "لون الباب"]

    def test_csv_utf8_sig(self):
        p = importer.parse(csv_bytes([self.HEAD, ["F2", "50035533", "2026-10-12", FAKE_ID, "أزرق"]]), "a.csv")
        r = p["rows"][0]
        self.assertEqual((r["unit_text"], r["permit_no"], r["end_date"]), ("F2", "50035533", "2026-10-12"))
        self.assertEqual(r["holder_id_last4"], "4321")
        self.assertEqual(r["extra_json"], {"لون الباب": "أزرق"})

    def test_csv_cp1256(self):
        p = importer.parse(csv_bytes([self.HEAD, ["الشهداء 9B", "1", "12/01/2027", "", ""]], "cp1256"), "w.csv")
        self.assertEqual(p["rows"][0]["unit_text"], "الشهداء 9B")
        self.assertEqual(p["rows"][0]["end_date"], "2027-01-12")

    def test_xlsx_header_on_row_3_with_a_link(self):
        data = xlsx_bytes([["م", "اسم الوحدة", "ملف التصريح", "تاريخ الإنتهاء (هجري)"],
                           [1, "F2", "f2.pdf", "01/05/1448"]],
                          link=((1, 2), "https://drive.google.com/file/d/abc"))
        p = importer.parse(data, "x.xlsx")
        self.assertEqual(p["header_row"], 3)
        r = p["rows"][0]
        self.assertEqual(r["doc_url"], "https://drive.google.com/file/d/abc")
        self.assertEqual(r["serial"], 1)
        if importer.dates.hijri_available():
            self.assertEqual(r["end_date"], "2026-10-12")

    def test_json_list(self):
        data = json.dumps([{"unit": "A5", "expiry": "2026-11-24", "permit no": "50036747"},
                           {"unit": "C08", "expiry": "2026-11-23", "permit no": "50036716"}]).encode()
        p = importer.parse(data, "x.json")
        self.assertEqual([r["unit_text"] for r in p["rows"]], ["A5", "C08"])
        self.assertEqual(p["rows"][1]["end_date"], "2026-11-23")

    def test_too_big_and_empty_are_refused(self):
        with self.assertRaises(importer.ImportError_):
            importer.parse(b"", "x.csv")
        with self.assertRaises(importer.ImportError_):
            importer.parse(b"x" * (importer.MAX_BYTES + 1), "x.csv")


class TestNothingDropped(DbCase):

    def test_row_count_in_equals_row_count_stored(self):
        rows = [["الشقة", "رقم التصريح", "تاريخ الانتهاء"]]
        for i in range(40):
            end = "2026-12-%02d" % (1 + i % 28) if i % 5 else "قريب"     # every 5th unreadable
            rows.append(["U%d" % i, "9%07d" % i, end])
        p = importer.parse(csv_bytes(rows), "many.csv")
        self.assertEqual(p["source_rows"], 40)
        res = importer.commit(p, "tester")
        self.assertEqual(res["inserted"], 40)
        self.assertEqual(res["needs_data"], 8)
        self.assertEqual(db.count_permits(), 40)
        nd = db.q("SELECT * FROM permits_permits WHERE needs_data=1")
        self.assertEqual(len(nd), 8)
        self.assertTrue(all(x["end_date"] is None and x["end_date_raw"] == "قريب" for x in nd))

    def test_blank_rows_are_the_only_rows_skipped(self):
        p = importer.parse(csv_bytes([["الشقة", "تاريخ الانتهاء"], ["A", "2026-12-01"], ["", ""], ["B", ""]]), "b.csv")
        self.assertEqual([r["unit_text"] for r in p["rows"]], ["A", "B"])
        self.assertEqual(p["rows"][1]["needs_data"], 1)

    def test_in_file_duplicates_are_both_kept_and_flagged(self):
        p = importer.parse(csv_bytes([["م", "الشقة", "رقم التصريح", "تاريخ الانتهاء"],
                                      ["19", "101a", "50037629", "2026-12-26"],
                                      ["44", "201b", "50037629", "2026-12-26"]]), "d.csv")
        res = importer.commit(p, "tester")
        self.assertEqual(res["inserted"], 2)
        stored = db.permits("active")
        self.assertEqual(len(stored), 2)
        for s in stored:
            codes = [i["code"] for i in json.loads(s["review_issues"])]
            self.assertEqual(codes, ["dup_permit_no"])
        self.assertIn("44", json.loads(stored[0]["review_issues"])[0]["text_ar"])

    def test_an_existing_number_is_skipped_unless_update_is_chosen(self):
        first = importer.parse(csv_bytes([["الشقة", "رقم التصريح", "تاريخ الانتهاء"], ["A", "77", "2026-12-01"]]), "a.csv")
        importer.commit(first, "t")
        again = importer.parse(csv_bytes([["الشقة", "رقم التصريح", "تاريخ الانتهاء"], ["A", "77", "2027-12-01"]]), "a.csv")
        importer.mark_existing(again)
        self.assertTrue(again["rows"][0]["exists_id"])
        self.assertEqual(importer.commit(again, "t")["skipped"], 1)
        self.assertEqual(db.permits("active")[0]["end_date"], "2026-12-01")
        res = importer.commit(again, "t", decisions={again["rows"][0]["row_index"]: "update"})
        self.assertEqual(res["updated"], 1)
        self.assertEqual(db.permits("active")[0]["end_date"], "2027-12-01")
        self.assertEqual(db.count_permits(), 1)


class TestReviewFixes(DbCase):

    def test_a_ten_digit_permit_number_is_kept_whole(self):
        cr = "1" + "0" * 5 + "3456"          # a commercial-registration-shaped number
        p = importer.parse(csv_bytes([["الشقة", "رقم التصريح", "تاريخ الانتهاء"], ["A", cr, "2026-12-01"]]), "x.csv")
        self.assertEqual(p["rows"][0]["permit_no"], cr)

    def test_a_manual_link_survives_an_update_import(self):
        importer.commit(importer.parse(csv_bytes([["الشقة", "رقم التصريح", "تاريخ الانتهاء"],
                                                  ["A", "77", "2026-12-01"]]), "a.csv"), "t")
        pid = db.permits("active")[0]["id"]
        db.update_permit(pid, {"listing_id": 5, "listing_link_kind": "manual"})
        pv = importer.from_onboarding(lambda: [{"id": 1, "unit_name": "A", "listing_id": 9,
                                                "license_no": "77", "license_expiry": "2027-12-01"}])
        importer.mark_existing(pv)
        importer.commit(pv, "t", decisions={pv["rows"][0]["row_index"]: "update"})
        p = db.permit(pid)
        self.assertEqual((p["listing_id"], p["listing_link_kind"], p["end_date"]), (5, "manual", "2027-12-01"))


class TestDatesAndColumns(DbCase):

    def test_ambiguous_is_flagged_in_the_preview(self):
        p = importer.parse(csv_bytes([["الشقة", "تاريخ الانتهاء"], ["A", "5/3/2027"]]), "x.csv")
        r = p["rows"][0]
        self.assertEqual(r["date_issue"], "ambiguous_day_month")
        self.assertEqual(r["preview"]["end_words"], "٥ مارس ٢٠٢٧")

    def test_guessed_column(self):
        p = importer.parse(csv_bytes([["الشقة", "عمود ١", "عمود ٢"], ["A", "2025-12-01", "2026-12-01"]]), "x.csv")
        r = p["rows"][0]
        self.assertEqual((r["start_date"], r["end_date"]), ("2025-12-01", "2026-12-01"))
        self.assertEqual(r["date_issue"], "guessed_column")

    def test_commit_reparses_and_ignores_a_tampered_client_date(self):
        """The routes layer re-runs parse() on the uploaded bytes; a preview row edited in
        the browser is never what gets stored. Here: the stored date is the file's."""
        raw = csv_bytes([["الشقة", "تاريخ الانتهاء"], ["A", "2026-12-01"]])
        tampered = importer.parse(raw, "x.csv")
        tampered["rows"][0]["end_date"] = "2099-01-01"
        fresh = importer.parse(raw, "x.csv")
        importer.commit(fresh, "t")
        self.assertEqual(db.permits("active")[0]["end_date"], "2026-12-01")


class TestPrivacy(DbCase):

    def test_id_column_keeps_last4_only_everywhere(self):
        p = importer.parse(csv_bytes([["الشقة", "السجل المدني", "ملاحظات", "تاريخ الانتهاء"],
                                      ["A", FAKE_ID, "هويته " + FAKE_ID, "2026-12-01"]]), "x.csv")
        importer.commit(p, "t")
        with closing(sqlite3.connect(bdb.db_path())) as cx:
            dump = "\n".join(cx.iterdump())
        self.assertNotIn(FAKE_ID, dump)
        self.assertNotIn(FAKE_ID, json.dumps(p, ensure_ascii=False))
        row = db.permits("active")[0]
        self.assertEqual(row["holder_id_last4"], "4321")
        self.assertIn("••••4321", row["notes"])


class TestOnboardingSourceIsReadOnly(DbCase):

    def test_no_writes_to_onb(self):
        with closing(bdb.connect()) as cx:
            cx.execute("CREATE TABLE onb_projects (id INTEGER PRIMARY KEY, unit_name TEXT, district TEXT,"
                       " listing_id INTEGER, license_no TEXT, license_expiry TEXT)")
            cx.execute("INSERT INTO onb_projects VALUES (1,'Hue 9','الملقا',11,'50043341','2027-05-07')")
            cx.execute("INSERT INTO onb_projects VALUES (2,'F9','العليا',12,'','')")
            cx.commit()

        def reader():
            with closing(bdb.connect()) as cx:
                return [dict(r) for r in cx.execute("SELECT * FROM onb_projects").fetchall()]

        with closing(sqlite3.connect(bdb.db_path())) as cx:
            before = cx.execute("SELECT * FROM onb_projects ORDER BY id").fetchall()
        p = importer.from_onboarding(reader)
        self.assertEqual(len(p["rows"]), 1)
        self.assertEqual(p["rows"][0]["permit_no"], "50043341")
        self.assertEqual(p["rows"][0]["permit_type"], engine.SEED_TYPE)
        importer.commit(p, "t", source="onboarding")
        with closing(sqlite3.connect(bdb.db_path())) as cx:
            after = cx.execute("SELECT * FROM onb_projects ORDER BY id").fetchall()
        self.assertEqual(before, after)
        self.assertEqual(db.permits("active")[0]["listing_id"], 11)

    def test_the_module_never_names_an_onb_write(self):
        src = open(importer.__file__, encoding="utf-8").read()
        self.assertNotRegex(src, r"(?i)(insert|update|delete)\s+(into\s+)?onb_")


if __name__ == "__main__":
    unittest.main()
